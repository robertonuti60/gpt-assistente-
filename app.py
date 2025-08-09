import os
import io
import json
import requests
from flask import Flask, request, jsonify
import msal

# --- opzionali per /read (PDF/DOCX) ---
# Li importiamo dentro le funzioni per non rompere il boot se mancano
# from pdfminer.high_level import extract_text
# from docx import Document

# ========= Config via variabili d’ambiente =========
CLIENT_ID     = os.environ.get("CLIENT_ID", "")
CLIENT_SECRET = os.environ.get("CLIENT_SECRET", "")
TENANT_ID     = os.environ.get("TENANT_ID", "")
ONEDRIVE_USER = os.environ.get("ONEDRIVE_USER", "")
API_KEY       = os.environ.get("API_KEY", "")
MAX_CHARS     = int(os.environ.get("MAX_CHARS", "500000"))

GRAPH_SCOPE = ["https://graph.microsoft.com/.default"]
GRAPH_BASE  = "https://graph.microsoft.com/v1.0"

app = Flask(__name__)


# ========= Guard API-key =========
def require_api_key(f):
    def wrapper(*args, **kwargs):
        if API_KEY and request.headers.get("x-api-key") != API_KEY:
            return jsonify({"error": "Forbidden"}), 403
        return f(*args, **kwargs)
    # mantieni il nome per Flask
    wrapper.__name__ = f.__name__
    return wrapper


# ========= Auth e chiamate Graph =========
def get_token():
    """
    Ritorna una access_token oppure solleva RuntimeError con info leggibili.
    """
    appc = msal.ConfidentialClientApplication(
        CLIENT_ID,
        authority=f"https://login.microsoftonline.com/{TENANT_ID}",
        client_credential=CLIENT_SECRET,
    )
    result = appc.acquire_token_silent(GRAPH_SCOPE, account=None) \
             or appc.acquire_token_for_client(scopes=GRAPH_SCOPE)

    if "access_token" not in result:
        # Normalizziamo l’errore per il client
        raise RuntimeError(json.dumps({
            "error": result.get("error"),
            "error_description": result.get("error_description"),
            "correlation_id": result.get("correlation_id"),
        }))
    return result["access_token"]


def gget(url, params=None, stream=False):
    """
    GET su Microsoft Graph con bearer; alza HTTPError se lo status non è 2xx
    """
    token = get_token()
    r = requests.get(url, headers={"Authorization": f"Bearer {token}"}, params=params, stream=stream, timeout=60)
    r.raise_for_status()
    return r


def path_url(path: str) -> str:
    """
    Costruisce la URL graf per un path OneDrive dell’utente
    """
    p = (path or "").lstrip("/")
    return f"{GRAPH_BASE}/users/{ONEDRIVE_USER}/drive/root:/{p}"


# ========= Endpoint base =========
@app.get("/")
def home():
    return {
        "name": "gpt-assistente-api",
        "ok": True,
        "endpoints": {
            "health": "/health",
            "search": "/search?q=term&path=<facoltativo>",
            "read": "POST /read { path: '/cartella/file.pdf' }",
            "debug_token": "/debug/token",
            "debug_drive": "/debug/drive",
            "debug_list": "/debug/list?path=<cartella>"
        },
        "note": "Usa l’intestazione 'x-api-key: <API_KEY>' per /search, /read e /debug/*"
    }


@app.get("/health")
def health():
    return {"ok": True}


@app.get("/search")
@require_api_key
def search():
    q    = request.args.get("q", "").strip()
    path = (request.args.get("path") or "").strip()

    if not q:
        return {"error": "Parametro 'q' obbligatorio"}, 400

    try:
        # Nel caso di path restringiamo la ricerca dentro la cartella
        if path:
            url = path_url(path) + f":/search(q='{q}')"
        else:
            url = f"{GRAPH_BASE}/users/{ONEDRIVE_USER}/drive/root/search(q='{q}')"

        data = gget(url).json()
        out = []
        for it in data.get("value", []):
            parent = it.get("parentReference", {}).get("path", "")
            # il path Graph ha "/drive/root:" all'inizio
            full   = (parent + "/" + it.get("name", "")).replace("/drive/root:", "")
            out.append({
                "name": it.get("name"),
                "path": full,
                "size": it.get("size"),
                "webUrl": it.get("webUrl"),
                "id": it.get("id"),
            })
        return {"results": out}
    except requests.HTTPError as e:
        try:
            return {"graph_error": e.response.json()}, 502
        except Exception:
            return {"graph_error": str(e)}, 502
    except Exception as e:
        # eventuale errore msal o altro
        return {"error": str(e)}, 500


@app.post("/read")
@require_api_key
def read():
    body = request.get_json(force=True, silent=True) or {}
    path = (body.get("path") or "").strip()
    if not path:
        return {"error": "Missing 'path'"}, 400

    try:
        # Scarica il binario del file
        content = gget(path_url(path) + ":/content", stream=True).content
        lower   = path.lower()

        if lower.endswith(".pdf"):
            from pdfminer.high_level import extract_text
            text = extract_text(io.BytesIO(content))
        elif lower.endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(content))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append("\t".join(c.text for c in row.cells))
            text = "\n".join(parts)
        elif lower.endswith(".txt"):
            text = content.decode("utf-8", errors="ignore")
        else:
            return {"error": "Unsupported file type"}, 415

        if len(text) > MAX_CHARS:
            text = text[:MAX_CHARS] + "\n\n[...troncato...]"

        return {"path": path, "text": text}
    except requests.HTTPError as e:
        try:
            return {"graph_error": e.response.json()}, 502
        except Exception:
            return {"graph_error": str(e)}, 502
    except Exception as e:
        return {"error": str(e)}, 500


# ========= Endpoint di DEBUG =========
@app.get("/debug/token")
@require_api_key
def debug_token():
    try:
        appc = msal.ConfidentialClientApplication(
            CLIENT_ID,
            authority=f"https://login.microsoftonline.com/{TENANT_ID}",
            client_credential=CLIENT_SECRET,
        )
        result = appc.acquire_token_silent(GRAPH_SCOPE, account=None) \
                 or appc.acquire_token_for_client(scopes=GRAPH_SCOPE)
        if "access_token" not in result:
            return {"msal_error": {
                "error": result.get("error"),
                "error_description": result.get("error_description"),
                "correlation_id": result.get("correlation_id")
            }}, 502
        return {"token": "OK", "scope": GRAPH_SCOPE, "tenant": TENANT_ID}
    except Exception as e:
        return {"exception": str(e)}, 500


@app.get("/debug/drive")
@require_api_key
def debug_drive():
    try:
        me = gget(f"{GRAPH_BASE}/users/{ONEDRIVE_USER}").json()
        drive = gget(f"{GRAPH_BASE}/users/{ONEDRIVE_USER}/drive").json()
        root = gget(f"{GRAPH_BASE}/users/{ONEDRIVE_USER}/drive/root/children?$top=10").json()
        return {
            "userPrincipalName": me.get("userPrincipalName"),
            "driveId": drive.get("id"),
            "rootItems": [i.get("name") for i in root.get("value", [])],
        }
    except requests.HTTPError as e:
        try:
            return {"http_error": e.response.json()}, 502
        except Exception:
            return {"http_error": str(e)}, 502
    except Exception as e:
        return {"error": str(e)}, 500


@app.get("/debug/list")
@require_api_key
def debug_list():
    p = (request.args.get("path") or "").strip().lstrip("/")
    url = f"{GRAPH_BASE}/users/{ONEDRIVE_USER}/drive/root/children" if not p \
          else f"{GRAPH_BASE}/users/{ONEDRIVE_USER}/drive/root:/{p}:/children"
    try:
        data = gget(url).json()
        return [i.get("name") for i in data.get("value", [])]
    except requests.HTTPError as e:
        try:
            return {"http_error": e.response.json()}, 502
        except Exception:
            return {"http_error": str(e)}, 502
    except Exception as e:
        return {"error": str(e)}, 500


# ========= Avvio locale =========
if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "8000")))
